import csv
import io
import json
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from sqlalchemy import select
from .models import Inspection, Revision, Dataset, User, Company, Audit, Anomaly, Deadline, now

ROME = ZoneInfo("Europe/Rome")
def quarter(year, q):
    if not 1 <= q <= 4 or not 2000 <= year <= 2200: raise ValueError("Trimestre non valido")
    return datetime(year, 3*q-2, 1, tzinfo=ROME), datetime(year+1 if q == 4 else year, 1 if q == 4 else 3*q+1, 1, tzinfo=ROME)
def display(s): return datetime.fromisoformat(s).astimezone(ROME).isoformat() if s else ""
def review_state(db, inspection_id, revision):
    records = db.scalars(select(Audit).where(Audit.kind == "review", Audit.target_id == inspection_id).order_by(Audit.at.desc())).all()
    return next((x.payload["state"] for x in records if x.payload["revision"] == revision), "NON_ESAMINATA")

def rows(db, area_id):
    result = []
    datasets = {}
    for i in db.scalars(select(Inspection).where(Inspection.area_id == area_id)):
        rev = db.get(Revision, (i.id, i.current_revision))
        p = rev.payload
        if i.dataset_id not in datasets: datasets[i.dataset_id] = {p["id"]: p for p in db.get(Dataset, i.dataset_id).payload["points"]}
        point = datasets[i.dataset_id][i.manhole_id]
        e = p["events"][-1]
        ev = rev.server_evaluation
        result.append({"inspection_id": i.id, "manhole_id": i.manhole_id, "collector": " | ".join(point["collectors"]), "manhole": point["code"],
                       "user": db.get(User, i.user_id).username, "company": db.get(Company, i.company_id).name,
                       "gps_at": e["acquired_at"], "completed_at": p["completed_at"], "received_at": i.received_at, "revision_received_at": rev.received_at,
                       "technical": p["sheet"], "status": p["status"], "distance_m": ev["distance_m"], "accuracy_m": e.get("accuracy_m"),
                       "gps": ev["state"], "exception": p["sheet"]["exception_reason"], "anomaly": p["sheet"]["anomaly_note"],
                       "review": review_state(db, i.id, i.current_revision), "revision": i.current_revision,
                       "deadline_id": p.get("deadline_id"), "synthetic": i.synthetic})
    return result

def snapshot(db, area_id, year, q, synthetic):
    start, end = quarter(year, q)
    all_rows = [r for r in rows(db, area_id) if r["synthetic"] == synthetic]
    selected = [r for r in all_rows if start <= datetime.fromisoformat(r["completed_at"]) < end]
    deadlines = []
    for d in db.scalars(select(Deadline).where(Deadline.area_id == area_id)):
        due = datetime.fromisoformat(d.due_at)
        visits = [r for r in all_rows if r["deadline_id"] == d.id and r["status"] == "COMPLETO"]
        completed = min((r["completed_at"] for r in visits), default=None)
        deadlines.append({"id": d.id, "manhole_id": d.manhole_id, "due_at": d.due_at, "completed_at": completed,
                          "in_period": start <= due < end, "on_time": completed is not None and datetime.fromisoformat(completed) <= due,
                          "late": completed is not None and datetime.fromisoformat(completed) > due,
                          "outstanding_at_period_end": due < end and (completed is None or datetime.fromisoformat(completed) >= end)})
    due_in_period = [d for d in deadlines if d["in_period"]]
    due_points = {d["manhole_id"] for d in due_in_period}
    covered_points = {d["manhole_id"] for d in due_in_period if d["completed_at"] is not None and datetime.fromisoformat(d["completed_at"]) < end}
    return {"period": f"{year}-T{q}", "version": "Identificativo immutabile dell'emissione", "extracted_at": now(), "timezone": "Europe/Rome",
            "synthetic": synthetic, "rows": selected, "deadlines": deadlines,
            "indicators": {"due": len(due_in_period), "on_time": sum(d["on_time"] for d in due_in_period),
                           "distinct_due": len(due_points), "distinct_due_covered": len(covered_points),
                           "coverage": len(covered_points)/len(due_points) if due_points else None,
                           "distinct_completed": len({r["manhole_id"] for r in selected if r["status"] == "COMPLETO"}),
                           "outstanding_at_period_end": sum(d["outstanding_at_period_end"] for d in deadlines)}}

HEADERS = ["period", "collector", "manhole", "user", "company", "gps_at", "completed_at", "received_at", "revision_received_at", "technical", "status", "distance_m", "accuracy_m", "gps", "exception", "anomaly", "review", "inspection_id", "revision", "synthetic", "extracted_at", "report_id"]
def tabular(s):
    out = []
    for row in s["rows"]:
        r = {**row, "period": s["period"], "extracted_at": display(s["extracted_at"]), "report_id": s["report_id"]}
        for key in ["gps_at", "completed_at", "received_at", "revision_received_at"]: r[key] = display(r[key])
        r["technical"] = json.dumps(r["technical"], ensure_ascii=False)
        out.append([r.get(h, "") for h in HEADERS])
    return out

def export_csv(s):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream, delimiter=";", quoting=csv.QUOTE_ALL)
    writer.writerow(HEADERS+["record_type"])
    metadata={"period":s["period"],"extracted_at":display(s["extracted_at"]),"report_id":s["report_id"],"synthetic":s["synthetic"]}
    writer.writerow([metadata.get(h,"") for h in HEADERS]+["EMISSIONE"])
    for row in [r+["CONTROLLO"] for r in tabular(s)]:
        # Spreadsheet formula injection mitigation. CSV has no text cell types.
        writer.writerow([("'"+v if v.lstrip().startswith(("=", "+", "-", "@")) or v.startswith(("\t", "\r", "\n")) else v) if isinstance(v, str) else v for v in row])
    return stream.getvalue().encode("utf-8-sig")

def export_xlsx(s):
    wb = Workbook()
    ws = wb.active
    ws.title = "Controlli"
    for row in [HEADERS] + tabular(s):
        ws.append(row)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str): cell.data_type, cell.number_format = "s", "@"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    meta = wb.create_sheet("Emissione")
    for k in ["report_id", "period", "extracted_at", "synthetic", "timezone"]: meta.append([k, str(s[k])])
    for k, v in s["indicators"].items(): meta.append([k, v])
    dl = wb.create_sheet("Scadenze")
    fields = ["id", "manhole_id", "due_at", "completed_at", "in_period", "on_time", "late", "outstanding_at_period_end"]
    dl.append(fields)
    for d in s["deadlines"]: dl.append([d[k] for k in fields])
    buff = io.BytesIO()
    wb.save(buff)
    return buff.getvalue()
